"""Generate ``registry.xlsx`` from the registry CSVs (PROJECT_CONTEXT §9). Never edit the xlsx by hand.

Sheets: Runs, Models, one leaderboard per showcase track, and Ablations (runs grouped by
tag without the seed suffix, using each run's best checkpoint: mean ± std).
"""
from __future__ import annotations

import statistics
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from racer.registry.registry import (LIBRARY_MAPS, MODEL_COLUMNS, RUN_COLUMNS, SHOWCASE, Registry, ablation_group,
                                     typed)

HEADER_FILL = PatternFill("solid", fgColor="1F2A37")
HEADER_FONT = Font(bold=True, color="FFFFFF")
GOOD_FILL = PatternFill("solid", fgColor="C6EFCE")
BEST_FONT = Font(bold=True, color="006100")
LINK_FONT = Font(color="0563C1", underline="single")
PCT_COLS = {"completion_rate", "crash_rate", "heldout_completion", "overall_completion", "score"}


def _sheet(wb: Workbook, title: str, columns: Sequence[str], rows: Sequence[Sequence[Any]]):
    ws = wb.create_sheet(title)
    ws.append(list(columns))
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(list(r))
    ws.freeze_panes = "B2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    for i, name in enumerate(columns, 1):
        width = max([len(str(name))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None][:200])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 48)
        if name in PCT_COLS:
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                row[0].number_format = "0%"
        elif name.startswith(("best_lap", "mean_lap", "lap_ratio", "mean_speed", "offroad", "mean_lat", "gap")):
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                row[0].number_format = "0.00"
    return ws


def _link(cell, target: str, text: str | None = None) -> None:
    cell.hyperlink = target
    cell.value = text or cell.value
    cell.font = LINK_FONT


def _thumb_prefix(m: dict) -> str:
    """Thumbnail file prefix of the protocol a registry row was evaluated with."""
    from racer.agents.evaluation import eval_filenames, protocol_version

    return eval_filenames(max(protocol_version(m.get("eval_protocol")), 1))[1]


def export(reg: Registry, out: Path | None = None) -> Path:
    out = Path(out) if out else reg.root / "registry.xlsx"
    runs = [{k: typed(v) for k, v in r.items()} for r in reg.runs()]
    models = [{k: typed(v) for k, v in r.items()} for r in reg.models()]
    wb = Workbook()
    wb.remove(wb.active)

    _sheet(wb, "Runs", RUN_COLUMNS, [[r.get(c) for c in RUN_COLUMNS] for r in runs])

    cols = MODEL_COLUMNS + ["thumbnail"]
    rows = sorted(models, key=lambda m: (str(m["run_id"]), m.get("steps") or 0))
    ws = _sheet(wb, "Models", cols, [[m.get(c) for c in MODEL_COLUMNS] + ["view"] for m in rows])
    path_col = MODEL_COLUMNS.index("path") + 1
    for i, m in enumerate(rows, 2):
        if m.get("path"):
            _link(ws.cell(i, path_col), f"../{m['path']}")
        if m.get("eval_date"):
            _link(ws.cell(i, len(cols)), f"../outputs/eval/{m['model_id']}/{_thumb_prefix(m)}{SHOWCASE[0]}.png")
        if m.get("is_best"):
            ws.cell(i, 1).font = BEST_FONT
    # Showcase leaderboards always; library maps (protocol v2) once some model has a lap there.
    lap_maps = [*SHOWCASE, *(t for t in LIBRARY_MAPS if any(m.get(f"best_lap_{t}") for m in models))]
    for t in lap_maps:     # green = the fastest lap on each map
        letter = get_column_letter(MODEL_COLUMNS.index(f"best_lap_{t}") + 1)
        if rows:
            rng = f"{letter}2:{letter}{len(rows) + 1}"
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f"AND({letter}2<>\"\",{letter}2=MIN(${letter}$2:${letter}${len(rows) + 1}))"], fill=GOOD_FILL))

    base = next((m for m in models if m["run_id"] == "BASE"), None)
    for t in lap_maps:
        lb = sorted((m for m in models if m.get(f"best_lap_{t}")), key=lambda m: m[f"best_lap_{t}"])
        ref = base.get(f"best_lap_{t}") if base else None
        lcols = ["rank", "model_id", "best_lap", "vs_baseline_s", "completion_rate", "heldout_completion",
                 "unseen_completion", "mean_speed", "is_best", "eval_protocol"]
        lrows = [[k, m["model_id"], m[f"best_lap_{t}"], (m[f"best_lap_{t}"] - ref) if ref else None,
                  m.get("completion_rate"), m.get("heldout_completion"), m.get("unseen_completion"),
                  m.get("mean_speed"), m.get("is_best"), m.get("eval_protocol")]
                 for k, m in enumerate(lb, 1)]
        ws = _sheet(wb, f"LB_{t}"[:31], lcols, lrows)
        for i, m in enumerate(lb, 2):
            _link(ws.cell(i, 2), f"../outputs/eval/{m['model_id']}/{_thumb_prefix(m)}{t}.png", m["model_id"])

    # Ablations: each run contributes its best checkpoint; group by tag minus seed suffix.
    best = {m["run_id"]: m for m in models if m.get("is_best")}
    groups: dict[str, list[tuple[dict, dict]]] = {}
    for r in runs:
        if r["run_id"] in best:
            groups.setdefault(ablation_group(str(r.get("tag") or r["run_id"])), []).append((r, best[r["run_id"]]))
    metrics = ["overall_completion", "unseen_completion", "heldout_completion", "test_completion", "completion_rate",
               "lap_ratio", "mean_speed", "offroad_pct"]
    acols = ["group", "n_runs", "runs", "physics", "reward", "obs_spec", "track_set"]
    for k in metrics:
        acols += [f"{k}_mean", f"{k}_std"]
    arows = []
    for g, items in sorted(groups.items()):
        r0 = items[0][0]
        row = [g, len(items), " ".join(r["run_id"] for r, _ in items), r0.get("physics_preset"),
               r0.get("reward_preset"), r0.get("obs_spec"), r0.get("track_set")]
        for k in metrics:
            vals = [m[k] for _, m in items if m.get(k) is not None]
            row += [statistics.fmean(vals) if vals else None, statistics.stdev(vals) if len(vals) > 1 else None]
        arows.append(row)
    ws = _sheet(wb, "Ablations", acols, arows)
    for i, name in enumerate(acols, 1):
        if name.startswith(("overall", "heldout", "completion")):
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                row[0].number_format = "0%"
        elif name.endswith(("_mean", "_std")):
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                row[0].number_format = "0.000"

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    print(export(Registry()))
