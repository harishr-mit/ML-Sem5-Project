"""Registry (T41): upsert, typed round trip, best-checkpoint ranking, xlsx export."""
import pytest
from openpyxl import load_workbook

from racer.registry.export_xlsx import export
from racer.registry.registry import Registry, ablation_group, typed


@pytest.fixture
def reg(tmp_path):
    r = Registry(tmp_path)
    r.upsert_run({"run_id": "R005", "tag": "e2_nogrip_s1", "physics_preset": "no_grip_limit", "total_steps": 100})
    for steps, comp, ratio in [(100, 0.5, 0.95), (200, 0.9, 1.10), (300, 0.9, 0.97), (400, None, None)]:
        row = {"model_id": f"R005_x@{steps}k", "run_id": "R005", "steps": steps, "path": f"models/R005_x@{steps}k"}
        if comp is not None:
            row.update(overall_completion=comp, lap_ratio=ratio, best_lap_training_ground=40 + steps / 100,
                       eval_date="2026-10-10T12:00:00")
        r.upsert_model(row)
    return r


def test_upsert_updates_in_place(reg):
    reg.upsert_model({"model_id": "R005_x@100k", "notes": "hello"})
    rows = reg.models()
    assert len(rows) == 4
    row = reg.model("R005_x@100k")
    assert row["notes"] == "hello" and row["steps"] == "100"      # other columns kept
    assert typed(row["overall_completion"]) == 0.5


def test_unknown_column_rejected(reg):
    with pytest.raises(KeyError):
        reg.upsert_model({"model_id": "x", "colour": "red"})


def test_best_is_highest_completion_then_fastest(reg):
    assert reg.set_best("R005") == "R005_x@300k"      # 0.9 completion, ratio 0.97 beats 1.10
    flags = {r["model_id"]: typed(r["is_best"]) for r in reg.models()}
    assert flags["R005_x@300k"] is True and flags["R005_x@200k"] is False


def test_export_has_all_sheets(reg, tmp_path):
    reg.set_best("R005")
    wb = load_workbook(export(reg, tmp_path / "r.xlsx"))
    assert wb.sheetnames == ["Runs", "Models", "LB_training_ground", "LB_technical", "LB_fast_flow", "Ablations"]
    assert wb["Ablations"]["A2"].value == "e2_nogrip"
    assert wb["LB_training_ground"]["B2"].value == "R005_x@100k"


def test_ablation_group_strips_seed():
    assert ablation_group("e4_rays9_s12") == "e4_rays9"
    assert ablation_group("ppo_rays5_tg") == "ppo_rays5_tg"


def test_backfill_skips_unfinished_exports(reg, tmp_path):
    """A GA checkpoint is written as .<id>.partial and renamed when complete; a leftover is never a model."""
    from racer.registry.registry import backfill
    models = tmp_path / "models"
    (models / ".R020_GA_x@5k.partial").mkdir(parents=True)
    (models / ".R020_GA_x@5k.partial" / "config.yaml").write_text("run_id: R020\n", encoding="utf-8")
    before = len(reg.models())
    assert backfill(reg, models) == []
    assert len(reg.models()) == before
